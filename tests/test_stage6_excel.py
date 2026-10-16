"""
ARL Stage 6 Test Suite: Excel Ingestion, Template-Preserving Output, and TQs.

Tests:
  1. Fuzzy column header matching and inspectable mapping reports.
  2. Manual column mapping overrides and optional LLM hook fallback.
  3. Multi-file Excel ingestion and signal/cable normalization.
  4. Data quality error codes E4001-E4010 on dirty data generating structured TQs.
  5. Template-preserving Excel output (preserves merged cells, formulas, styles).
  6. Atomic writing behavior and safe workbook reopening.
  7. Exit Gate: Complete single-cable workflow from Excel MDB to validated output workbook.
"""

from pathlib import Path
import tempfile
import openpyxl
import pandas as pd
import pytest

from arl import (
    AllocationResult,
    ExcelWriter,
    GreedySolver,
    Placement,
    TQManager,
    Validator,
    check_data_quality,
    create_default_template,
    read_excel_mdb,
    resolve_column_mapping,
    run_allocation_pipeline,
)


@pytest.fixture
def tmp_dir():
    with tempfile.TemporaryDirectory() as td:
        yield Path(td)


def _create_sample_excel(
    file_path: Path,
    rows: list,
    columns: list = None,
    sheet_name: str = "IO_MDB",
):
    """Helper to write sample rows to Excel."""
    if columns is None:
        columns = ["Tag Name", "I/O Type", "Cable No", "Pair #", "Channel Seg", "Plant Area"]
    df = pd.DataFrame(rows, columns=columns)
    df.to_excel(file_path, sheet_name=sheet_name, index=False)
    return file_path


class TestExcelIngestion:
    """Test Excel ingestion, column resolution, and normalization."""

    def test_fuzzy_matching_and_inspectable_mapping(self, tmp_dir):
        """Fuzzy column matcher matches non-standard headers and produces inspectable report."""
        cols = ["Instrument Tag", "I/O", "Multicore Cable", "Pair No.", "Segregation", "Unit Area"]
        mapping, details, unmapped = resolve_column_mapping(cols)

        assert mapping["Instrument Tag"] == "tag"
        assert mapping["I/O"] == "io_type"
        assert mapping["Multicore Cable"] == "cable"
        assert mapping["Pair No."] == "pair"
        assert mapping["Segregation"] == "seg"
        assert mapping["Unit Area"] == "area"
        assert len(unmapped) == 0

        # Verify mapping details report
        assert any(d.raw_column == "Instrument Tag" and d.canonical_field == "tag" for d in details)

    def test_manual_mapping_overrides(self, tmp_dir):
        """User manual mapping overrides fuzzy matching."""
        cols = ["SpecialColumnX", "Channel"]
        manual = {"SpecialColumnX": "tag", "Channel": "io_type"}
        mapping, details, unmapped = resolve_column_mapping(cols, manual_overrides=manual)

        assert mapping["SpecialColumnX"] == "tag"
        assert mapping["Channel"] == "io_type"
        assert any(d.method == "manual" for d in details)

    def test_multi_file_ingestion(self, tmp_dir):
        """Ingest signals distributed across multiple Excel workbooks."""
        f1 = tmp_dir / "mdb_part1.xlsx"
        f2 = tmp_dir / "mdb_part2.xlsx"

        _create_sample_excel(
            f1,
            [
                ["10-FT-001", "AI", "CBL-01", 1, "A", "AREA_01"],
                ["10-FT-002", "AI", "CBL-01", 2, "A", "AREA_01"],
            ],
        )
        _create_sample_excel(
            f2,
            [
                ["10-FT-003", "AO", "CBL-02", 1, "B", "AREA_01"],
                ["10-FT-004", "AO", "CBL-02", 2, "B", "AREA_01"],
            ],
        )

        project, tqs, col, report = read_excel_mdb([f1, f2])

        assert len(project.get("signals", [])) == 4
        assert len(project.get("cables", [])) == 2
        assert report.signals_created == 4
        assert len(report.files_processed) == 2


class TestDataQualityAndTQs:
    """Test data quality checks E4001-E4010 on dirty workbooks."""

    def test_missing_required_field_e4001(self, tmp_dir):
        f = tmp_dir / "dirty_e4001.xlsx"
        _create_sample_excel(f, [["", "AI", "CBL-01", 1, "A", "AREA_01"]])  # Empty tag
        project, tqs, col, report = read_excel_mdb(f)
        assert col.has_errors
        assert any(e.code == "E4001" for e in col.errors)

    def test_duplicate_tag_e4009(self, tmp_dir):
        f = tmp_dir / "dirty_e4009.xlsx"
        _create_sample_excel(
            f,
            [
                ["PT-100", "AI", "CBL-01", 1, "A", "AREA_01"],
                ["PT-100", "AI", "CBL-01", 2, "A", "AREA_01"],  # Duplicate tag
            ],
        )
        project, tqs, col, report = read_excel_mdb(f)
        assert any(t.code == "E4009" for t in tqs.get_all())

    def test_duplicate_pair_index_e4004(self, tmp_dir):
        f = tmp_dir / "dirty_e4004.xlsx"
        _create_sample_excel(
            f,
            [
                ["PT-101", "AI", "CBL-01", 1, "A", "AREA_01"],
                ["PT-102", "AI", "CBL-01", 1, "A", "AREA_01"],  # Duplicate pair 1 in CBL-01
            ],
        )
        project, tqs, col, report = read_excel_mdb(f)
        assert any(t.code == "E4004" for t in tqs.get_all())

    def test_invalid_io_type_e4006(self, tmp_dir):
        f = tmp_dir / "dirty_e4006.xlsx"
        _create_sample_excel(
            f,
            [["PT-103", "UNKNOWN_PROTOCOL_IO", "CBL-01", 1, "A", "AREA_01"]],
        )
        project, tqs, col, report = read_excel_mdb(f)
        assert any(t.code == "E4006" for t in tqs.get_all())

    def test_tq_schema_fields(self, tmp_dir):
        """Verify TQs generated adhere to Part 6.4 schema."""
        f = tmp_dir / "tq_schema_check.xlsx"
        _create_sample_excel(
            f,
            [
                ["TT-01", "AI", "CBL-01", 1, "A", "AREA_01"],
                ["TT-01", "AI", "CBL-01", 2, "A", "AREA_01"],
            ],
        )
        project, tqs, col, report = read_excel_mdb(f)
        all_tqs = tqs.get_all()
        assert len(all_tqs) > 0
        tq = all_tqs[0]

        # Part 6.4 required fields
        assert tq.tq_id.startswith("TQ-")
        assert tq.severity in ("critical", "high", "medium", "low")
        assert tq.code.startswith("E")
        assert tq.title != ""
        assert "signals" in tq.affected
        assert "cables" in tq.affected
        assert "pairs" in tq.affected
        assert "unplaced_signals" in tq.impact
        assert isinstance(tq.suggestions, list)
        assert isinstance(tq.auto_fix_available, bool)

    def test_unknown_cable_size_e4002(self, tmp_dir):
        """E4002 flagged when cable size is non-standard (e.g. 7 pairs)."""
        f = tmp_dir / "dirty_e4002.xlsx"
        cols = ["Tag Name", "I/O Type", "Cable No", "Pair #", "Cable Size"]
        _create_sample_excel(
            f,
            [["TT-10", "AI", "CBL-NONSTD", 1, 7]],
            columns=cols,
        )
        project, tqs, col, report = read_excel_mdb(f)
        assert any(e.code == "E4002" for e in col.errors)

    def test_pair_outside_cable_size_e4010(self, tmp_dir):
        """E4010 flagged when pair index exceeds cable size."""
        f = tmp_dir / "dirty_e4010.xlsx"
        cols = ["Tag Name", "I/O Type", "Cable No", "Pair #", "Cable Size"]
        _create_sample_excel(
            f,
            [["TT-11", "AI", "CBL-SMALL", 10, 6]],  # Pair 10 in size 6 cable
            columns=cols,
        )
        project, tqs, col, report = read_excel_mdb(f)
        assert any(t.code == "E4010" for t in tqs.get_all())

    def test_group_seg_mismatch_e4005(self, tmp_dir):
        """E4005 flagged when signal group contains both Seg A and Seg B signals."""
        f = tmp_dir / "dirty_e4005.xlsx"
        cols = ["Tag Name", "I/O Type", "Cable No", "Pair #", "Channel Seg", "Group"]
        _create_sample_excel(
            f,
            [
                ["SG-01", "AI", "CBL-01", 1, "A", "GRP-MIX"],
                ["SG-02", "AI", "CBL-01", 2, "B", "GRP-MIX"],  # Seg B in same group as Seg A
            ],
            columns=cols,
        )
        project, tqs, col, report = read_excel_mdb(f)
        assert any(t.code == "E4005" for t in tqs.get_all())

    def test_llm_mapper_hook(self, tmp_dir):
        """Optional LLM mapper hook is called when use_llm=True."""
        cols = ["CustomTagCol", "CustomIOCol"]
        called_with = []

        def mock_llm_mapper(unmapped_cols):
            called_with.extend(unmapped_cols)
            return {"CustomTagCol": "tag", "CustomIOCol": "io_type"}

        mapping, details, unmapped = resolve_column_mapping(
            cols,
            use_llm=True,
            llm_mapper=mock_llm_mapper,
        )
        assert called_with == ["CustomTagCol", "CustomIOCol"]
        assert mapping["CustomTagCol"] == "tag"
        assert mapping["CustomIOCol"] == "io_type"
        assert any(d.method == "llm" for d in details)



class TestTemplatePreservingOutput:
    """Test template preservation, atomic writing, and multi-sheet generation."""

    def test_template_preservation(self, tmp_dir):
        """Template formulas, styles, and merged cells are preserved."""
        tmpl_file = tmp_dir / "client_template.xlsx"
        out_file = tmp_dir / "allocation_output.xlsx"

        # 1. Create a custom template with merged cells, formula, and styling
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Allocation"
        ws.merge_cells("A1:G1")
        ws["A1"] = "CLIENT CUSTOM DCS TEMPLATE BANNER"
        ws["A1"].font = openpyxl.styles.Font(size=16, bold=True, color="FF0000")
        ws["K10"] = "FormulaCheck:"
        ws["L10"] = "=SUM(10, 20)"  # Formula to preserve
        wb.save(tmpl_file)

        # 2. Run greedy solver on a sample project
        from arl import Baseplate, Cabinet, CIOC, Project, Signal
        bp = Baseplate(id="BP-01", capacity=12)
        cioc = CIOC(id="CIOC-01", baseplates=[bp])
        cab = Cabinet(id="CAB-01", ciocs=[cioc])
        sig = Signal(tag="10-PT-100", io_type="AI", cable="CBL-01", pair=1)
        proj = Project(name="TestProj", cabinets=[cab], signals=[sig])

        alloc_result = GreedySolver().solve(proj)
        val_result = Validator().validate(proj, alloc_result)

        # 3. Write using template
        writer = ExcelWriter()
        writer.write_allocation_workbook(
            output_path=out_file,
            allocation_result=alloc_result,
            project=proj,
            template_path=tmpl_file,
            validation_result=val_result,
        )

        assert out_file.exists()

        # 4. Reopen and verify preserved elements
        reopened = openpyxl.load_workbook(out_file)
        alloc_ws = reopened["Allocation"]

        # Check merged cell preserved
        assert "A1:G1" in [str(m) for m in alloc_ws.merged_cells.ranges]
        assert alloc_ws["A1"].value == "CLIENT CUSTOM DCS TEMPLATE BANNER"
        assert alloc_ws["L10"].value == "=SUM(10, 20)"  # Formula preserved!

        # Check all 5 required sheets exist
        for s in ["Allocation", "TQ_List", "Rule_Check", "Summary", "Impact"]:
            assert s in reopened.sheetnames

        # Check allocation row columns
        # Row 3 should be the placed signal
        assert alloc_ws.cell(row=3, column=1).value == "10-PT-100"
        assert alloc_ws.cell(row=3, column=2).value == "CAB-01"
        assert alloc_ws.cell(row=3, column=3).value == "CIOC-01"
        assert alloc_ws.cell(row=3, column=4).value == "BP-01"
        assert alloc_ws.cell(row=3, column=5).value == 1  # Slot 1
        assert alloc_ws.cell(row=3, column=6).value == "AI"  # CHARM Type
        assert alloc_ws.cell(row=3, column=7).value == "STANDARD"  # TB Type

    def test_atomic_write_safety(self, tmp_dir):
        """Atomic write ensures valid destination and clean temp removal."""
        out_file = tmp_dir / "atomic_target.xlsx"
        from arl import Baseplate, Cabinet, CIOC, Project, Signal
        bp = Baseplate(id="BP-01", capacity=12)
        cioc = CIOC(id="CIOC-01", baseplates=[bp])
        cab = Cabinet(id="CAB-01", ciocs=[cioc])
        sig = Signal(tag="SIG-01", io_type="DI", cable="CBL-01", pair=1)
        proj = Project(name="AtomicTest", cabinets=[cab], signals=[sig])

        alloc_result = GreedySolver().solve(proj)
        writer = ExcelWriter()
        writer.write_allocation_workbook(out_file, alloc_result, proj)

        assert out_file.exists()
        # Verify no temporary files left in directory
        temp_files = [f for f in tmp_dir.iterdir() if f.name.startswith(".tmp_")]
        assert len(temp_files) == 0


class TestExitGateSingleCableWorkflow:
    """Exit Gate requirement: End-to-end single-cable allocation workflow."""

    def test_single_cable_end_to_end(self, tmp_dir):
        """
        Complete a single-cable workflow from Excel input to normalized project data,
        greedy allocation, validation, TQ reporting, and filled output template.
        """
        mdb_file = tmp_dir / "SingleCable_MDB.xlsx"
        template_file = tmp_dir / "DeltaV_Template.xlsx"
        output_file = tmp_dir / "SingleCable_Allocated.xlsx"

        # 1. Create a 6-pair single-cable MDB with 5 active signals and 1 wired spare
        cable_rows = [
            ["10-FT-101", "AI", "CBL-TRUNK-01", 1, "A", "AREA_01"],
            ["10-PT-102", "AI", "CBL-TRUNK-01", 2, "A", "AREA_01"],
            ["10-TT-103", "AI", "CBL-TRUNK-01", 3, "A", "AREA_01"],
            ["10-XV-104", "DO_SOV", "CBL-TRUNK-01", 4, "A", "AREA_01"],
            ["10-ZSO-105", "DI", "CBL-TRUNK-01", 5, "A", "AREA_01"],
            # Pair 6 left unassigned as spare
        ]
        _create_sample_excel(mdb_file, cable_rows)

        # 2. Create the base template
        create_default_template(template_file)
        assert template_file.exists()

        # 3. Execute the full allocation pipeline
        pipeline_res = run_allocation_pipeline(
            mdb_paths=mdb_file,
            output_path=output_file,
            template_path=template_file,
        )

        # 4. Verify exit criteria
        assert pipeline_res.success is True
        assert pipeline_res.total_signals == 5
        assert pipeline_res.placed_signals == 5
        assert pipeline_res.unplaced_signals == 0
        assert pipeline_res.validation_passed is True
        assert output_file.exists()

        # 5. Reopen output file and verify content
        wb = openpyxl.load_workbook(output_file)
        ws_alloc = wb["Allocation"]

        tags_placed = [ws_alloc.cell(row=r, column=1).value for r in range(3, 8)]
        assert tags_placed == ["10-FT-101", "10-PT-102", "10-TT-103", "10-XV-104", "10-ZSO-105"]

        # Verify XV-104 has correct hardware assignment for DO_SOV
        xv_row = next(r for r in range(3, 8) if ws_alloc.cell(row=r, column=1).value == "10-XV-104")
        assert ws_alloc.cell(row=xv_row, column=6).value == "HIGH_SIDE_DO"
        assert ws_alloc.cell(row=xv_row, column=7).value == "FUSED_INJECTED_POWER"

        # Verify summary sheet KPIs
        ws_sum = wb["Summary"]
        assert ws_sum["C3"].value == 5  # Total Signals
        assert ws_sum["C4"].value == 5  # Placed Signals
        assert ws_sum["C12"].value == "VALID"  # Independent Validation Status
