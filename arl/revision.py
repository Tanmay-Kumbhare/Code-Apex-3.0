"""
ARL Revision Comparison and Move Minimization Engine.

Implements Part 10 of the ARL v2 specification:
  - 10.1 Diff:
      - Diff Rev A vs Rev B MDB (added, deleted, modified signals & cables).
      - Diff rulesets (added, removed, modified, overridden rules).
  - 10.2 Pinning:
      - Unchanged signals pinned to Rev A slots.
      - Pins are soft by default; hard pins only if group and cable unchanged.
      - Relaxation order: ungrouped -> soft signals -> group-changed -> safety last.
      - Move budget enforcement: exceeding raises E6003.
      - Error handling: E6001 (invalid pinned slot), E6002 (new rule invalidates pins),
        E6003 (move budget exceeded), E6004 (Rev A allocation missing).
  - 10.3 Impact Report:
      - Summary metrics (added, deleted, moved, unchanged counts).
      - Cabinets hit, move count, rule changes, and revision TQs.
"""

from __future__ import annotations

from enum import Enum, IntEnum
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple, Union
import openpyxl
from pydantic import BaseModel, ConfigDict, Field

from arl.ast import BinaryOp, ForbidEffect, Override, RequireEffect, Rule, Ruleset
from arl.entities import Baseplate, Cabinet, Cable, CIOC, Entity, Group, Project, Signal
from arl.errors import ArlError, ErrorCollector, Phase, Severity, make_error
from arl.excel_ingest import IngestionReport, read_excel_mdb
from arl.excel_output import ExcelWriter
from arl.schema import Schema, get_base_schema
from arl.solver_greedy import AllocationResult, Placement
from arl.tq import TQ, TQManager
from arl.validate import ValidationResult, Validator


class SignalChangeType(str, Enum):
    ADDED = "added"
    DELETED = "deleted"
    MODIFIED = "modified"
    UNCHANGED = "unchanged"


class SignalDiff(BaseModel):
    """Difference details for a single instrument signal."""
    tag: str
    change_type: SignalChangeType
    old_signal: Optional[Dict[str, Any]] = None
    new_signal: Optional[Dict[str, Any]] = None
    field_changes: Dict[str, Tuple[Any, Any]] = Field(default_factory=dict)

    model_config = ConfigDict(arbitrary_types_allowed=True)


class MDBDiff(BaseModel):
    """Diff between Rev A and Rev B project MDBs."""
    added: List[str] = Field(default_factory=list)        # Tags
    deleted: List[str] = Field(default_factory=list)      # Tags
    modified: List[SignalDiff] = Field(default_factory=list)
    unchanged: List[str] = Field(default_factory=list)    # Tags

    @property
    def added_count(self) -> int:
        return len(self.added)

    @property
    def deleted_count(self) -> int:
        return len(self.deleted)

    @property
    def modified_count(self) -> int:
        return len(self.modified)

    @property
    def unchanged_count(self) -> int:
        return len(self.unchanged)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "added_count": self.added_count,
            "deleted_count": self.deleted_count,
            "modified_count": self.modified_count,
            "unchanged_count": self.unchanged_count,
            "added": self.added,
            "deleted": self.deleted,
            "modified": [m.model_dump() for m in self.modified],
            "unchanged": self.unchanged,
        }


class RuleChangeType(str, Enum):
    ADDED = "added"
    REMOVED = "removed"
    MODIFIED = "modified"
    OVERRIDDEN = "overridden"


class RuleDiff(BaseModel):
    """Difference in a single DSL rule."""
    name: str
    rule_id: Optional[str] = None
    change_type: RuleChangeType
    details: str = ""

    model_config = ConfigDict(arbitrary_types_allowed=True)


class RulesetDiff(BaseModel):
    """Diff between ruleset Rev A and Rev B."""
    added_rules: List[RuleDiff] = Field(default_factory=list)
    removed_rules: List[RuleDiff] = Field(default_factory=list)
    modified_rules: List[RuleDiff] = Field(default_factory=list)
    overridden_rules: List[RuleDiff] = Field(default_factory=list)

    @property
    def total_changes(self) -> int:
        return (
            len(self.added_rules)
            + len(self.removed_rules)
            + len(self.modified_rules)
            + len(self.overridden_rules)
        )

    def to_dict(self) -> Dict[str, Any]:
        return {
            "total_changes": self.total_changes,
            "added_count": len(self.added_rules),
            "removed_count": len(self.removed_rules),
            "modified_count": len(self.modified_rules),
            "overridden_count": len(self.overridden_rules),
            "added_rules": [r.model_dump() for r in self.added_rules],
            "removed_rules": [r.model_dump() for r in self.removed_rules],
            "modified_rules": [r.model_dump() for r in self.modified_rules],
            "overridden_rules": [r.model_dump() for r in self.overridden_rules],
        }


class PinKind(str, Enum):
    HARD = "hard"
    SOFT = "soft"


class RelaxationPriority(IntEnum):
    """
    Priority order for pin relaxation as specified in Part 10.2:
      ungrouped -> soft signals -> group-changed -> safety last
    Lower numerical values are relaxed first.
    """
    UNGROUPED = 1
    SOFT_SIGNAL = 2
    GROUP_CHANGED = 3
    SAFETY_LAST = 4
    HARD_PIN = 5


class SignalPin(BaseModel):
    """Placement pin pinning a signal to its previous Rev A hardware location."""
    tag: str
    kind: PinKind
    cabinet_id: str
    cioc_id: str
    baseplate_id: str
    slot: int
    priority: RelaxationPriority
    is_safety: bool = False
    is_relaxed: bool = False
    relaxation_reason: Optional[str] = None

    model_config = ConfigDict(arbitrary_types_allowed=True)


class MoveRecord(BaseModel):
    """Record of a signal moving between revisions."""
    tag: str
    old_cabinet: str
    old_cioc: str
    old_baseplate: str
    old_slot: int
    new_cabinet: str
    new_cioc: str
    new_baseplate: str
    new_slot: int
    reason: str = "Re-allocated"

    model_config = ConfigDict(arbitrary_types_allowed=True)


class RevisionResult(BaseModel):
    """Complete outcome of a revision comparison and allocation run."""
    success: bool
    allocation_result: AllocationResult
    mdb_diff: MDBDiff
    ruleset_diff: Optional[RulesetDiff] = None
    moved_signals: List[MoveRecord] = Field(default_factory=list)
    move_count: int = 0
    max_moves: Optional[int] = None
    budget_exceeded: bool = False
    cabinets_hit: List[str] = Field(default_factory=list)
    tqs: List[TQ] = Field(default_factory=list)
    impact_summary: Dict[str, Any] = Field(default_factory=dict)
    output_path: Optional[str] = None

    model_config = ConfigDict(arbitrary_types_allowed=True)


# =========================================================================
# Diff Functions (10.1)
# =========================================================================

def diff_mdb(project_a: Project, project_b: Project) -> MDBDiff:
    """
    Diff signals between Rev A and Rev B project models.
    Identifies added, deleted, modified, and unchanged signals.
    """
    sigs_a = {str(s.get("tag")): s for s in project_a.get("signals", []) if s.get("tag")}
    sigs_b = {str(s.get("tag")): s for s in project_b.get("signals", []) if s.get("tag")}

    tags_a = set(sigs_a.keys())
    tags_b = set(sigs_b.keys())

    added = sorted(list(tags_b - tags_a))
    deleted = sorted(list(tags_a - tags_b))

    common_tags = sorted(list(tags_a & tags_b))
    modified: List[SignalDiff] = []
    unchanged: List[str] = []

    fields_to_compare = [
        "io_type", "cable", "pair", "group", "area", "unit",
        "seg", "safety", "voltage", "redundancy", "desc",
    ]

    for tag in common_tags:
        sig_a = sigs_a[tag]
        sig_b = sigs_b[tag]

        field_changes: Dict[str, Tuple[Any, Any]] = {}
        for f in fields_to_compare:
            val_a = sig_a.get(f)
            val_b = sig_b.get(f)
            if val_a != val_b:
                field_changes[f] = (val_a, val_b)

        if field_changes:
            modified.append(
                SignalDiff(
                    tag=tag,
                    change_type=SignalChangeType.MODIFIED,
                    old_signal=dict(sig_a),
                    new_signal=dict(sig_b),
                    field_changes=field_changes,
                )
            )
        else:
            unchanged.append(tag)

    return MDBDiff(
        added=added,
        deleted=deleted,
        modified=modified,
        unchanged=unchanged,
    )


def diff_rulesets(
    ruleset_a: Optional[Ruleset],
    ruleset_b: Optional[Ruleset],
) -> RulesetDiff:
    """
    Diff rules between ruleset Rev A and ruleset Rev B.
    Identifies added, removed, modified, and overridden rules.
    """
    if ruleset_a is None and ruleset_b is None:
        return RulesetDiff()

    rules_a = {r.name: r for r in (ruleset_a.rules if ruleset_a else [])}
    rules_b = {r.name: r for r in (ruleset_b.rules if ruleset_b else [])}

    added: List[RuleDiff] = []
    removed: List[RuleDiff] = []
    modified: List[RuleDiff] = []
    overridden: List[RuleDiff] = []

    # Detect added
    for name, rb in rules_b.items():
        if name not in rules_a:
            added.append(RuleDiff(name=name, rule_id=rb.id, change_type=RuleChangeType.ADDED, details="Rule added in Rev B"))

    # Detect removed
    for name, ra in rules_a.items():
        if name not in rules_b:
            removed.append(RuleDiff(name=name, rule_id=ra.id, change_type=RuleChangeType.REMOVED, details="Rule removed in Rev B"))

    # Detect modified
    for name in sorted(set(rules_a.keys()) & set(rules_b.keys())):
        ra = rules_a[name]
        rb = rules_b[name]

        diffs = []
        if ra.priority != rb.priority:
            diffs.append(f"priority changed: {ra.priority} -> {rb.priority}")
        if ra.weight != rb.weight:
            diffs.append(f"weight changed: {ra.weight} -> {rb.weight}")
        if str(ra.when) != str(rb.when):
            diffs.append("condition ('when') modified")
        if len(ra.effects) != len(rb.effects) or str(ra.effects) != str(rb.effects):
            diffs.append("rule effects modified")

        if diffs:
            modified.append(RuleDiff(name=name, rule_id=rb.id or ra.id, change_type=RuleChangeType.MODIFIED, details="; ".join(diffs)))

    # Detect overridden in Rev B
    if ruleset_b and ruleset_b.overrides:
        for ovr in ruleset_b.overrides:
            target_name = ovr.target_rule
            overridden.append(
                RuleDiff(
                    name=target_name,
                    change_type=RuleChangeType.OVERRIDDEN,
                    details=f"Overridden in Rev B (target: {target_name})",
                )
            )

    return RulesetDiff(
        added_rules=added,
        removed_rules=removed,
        modified_rules=modified,
        overridden_rules=overridden,
    )


def diff_allocations(
    alloc_a: Dict[str, Placement],
    alloc_b: Dict[str, Placement],
) -> Dict[str, Any]:
    """
    Diff signal allocations between Rev A and Rev B placements.
    """
    tags_a = set(alloc_a.keys())
    tags_b = set(alloc_b.keys())

    common = tags_a & tags_b
    moved: List[MoveRecord] = []
    unchanged: List[str] = []
    cabinets_hit: Set[str] = set()

    for tag in common:
        pa = alloc_a[tag]
        pb = alloc_b[tag]

        if (
            pa.cabinet_id != pb.cabinet_id
            or pa.cioc_id != pb.cioc_id
            or pa.baseplate_id != pb.baseplate_id
            or pa.slot != pb.slot
        ):
            moved.append(
                MoveRecord(
                    tag=tag,
                    old_cabinet=pa.cabinet_id,
                    old_cioc=pa.cioc_id,
                    old_baseplate=pa.baseplate_id,
                    old_slot=pa.slot,
                    new_cabinet=pb.cabinet_id,
                    new_cioc=pb.cioc_id,
                    new_baseplate=pb.baseplate_id,
                    new_slot=pb.slot,
                    reason="Slot or enclosure relocated",
                )
            )
            cabinets_hit.add(pa.cabinet_id)
            cabinets_hit.add(pb.cabinet_id)
        else:
            unchanged.append(tag)

    for tag in tags_b - tags_a:
        cabinets_hit.add(alloc_b[tag].cabinet_id)

    for tag in tags_a - tags_b:
        cabinets_hit.add(alloc_a[tag].cabinet_id)

    return {
        "moved": moved,
        "unchanged": unchanged,
        "moved_count": len(moved),
        "unchanged_count": len(unchanged),
        "cabinets_hit": sorted(list(cabinets_hit)),
    }


# =========================================================================
# Pinning Engine & Move Minimization (10.2)
# =========================================================================

class PinManager:
    """
    Manages generation, validation, relaxation, and move tracking for signal pins.
    """

    @staticmethod
    def create_pins(
        mdb_diff: MDBDiff,
        project_b: Project,
        alloc_a: Dict[str, Placement],
    ) -> Dict[str, SignalPin]:
        """
        Generate pins for Rev B signals based on Rev A placements and MDB diff:
          - Unchanged signals are pinned to their Rev A slots.
          - Pins are soft by default; hard pins only if group and cable unchanged.
          - Relaxation priority order: ungrouped -> soft signals -> group-changed -> safety last.
        """
        pins: Dict[str, SignalPin] = {}
        sigs_b = {str(s.get("tag")): s for s in project_b.get("signals", []) if s.get("tag")}
        modified_by_tag = {m.tag: m for m in mdb_diff.modified}

        for tag, pl_a in alloc_a.items():
            if tag not in sigs_b:
                # Deleted signal in Rev B
                continue

            sig_b = sigs_b[tag]
            is_safety = bool(sig_b.get("safety", False))
            group = sig_b.get("group")
            is_ungrouped = not group or group in ("", "None", "_ungrouped_")

            if tag in mdb_diff.unchanged:
                # Completely unchanged: group and cable are guaranteed unchanged
                kind = PinKind.HARD
                priority = RelaxationPriority.HARD_PIN
            elif tag in modified_by_tag:
                diff_rec = modified_by_tag[tag]
                group_changed = "group" in diff_rec.field_changes
                cable_changed = "cable" in diff_rec.field_changes

                if not group_changed and not cable_changed:
                    # Group and cable unchanged! Eligible for hard pin if no hardware change
                    io_type_changed = "io_type" in diff_rec.field_changes
                    if not io_type_changed:
                        kind = PinKind.HARD
                        priority = RelaxationPriority.HARD_PIN
                    else:
                        kind = PinKind.SOFT
                        priority = RelaxationPriority.SOFT_SIGNAL
                elif group_changed:
                    kind = PinKind.SOFT
                    priority = RelaxationPriority.GROUP_CHANGED
                elif is_ungrouped:
                    kind = PinKind.SOFT
                    priority = RelaxationPriority.UNGROUPED
                else:
                    kind = PinKind.SOFT
                    priority = RelaxationPriority.SOFT_SIGNAL
            else:
                kind = PinKind.SOFT
                priority = RelaxationPriority.SOFT_SIGNAL

            # Safety signals have highest retention priority: relaxed last
            if is_safety:
                priority = RelaxationPriority.SAFETY_LAST

            pins[tag] = SignalPin(
                tag=tag,
                kind=kind,
                cabinet_id=pl_a.cabinet_id,
                cioc_id=pl_a.cioc_id,
                baseplate_id=pl_a.baseplate_id,
                slot=pl_a.slot,
                priority=priority,
                is_safety=is_safety,
            )

        return pins

    @staticmethod
    def validate_pins(
        pins: Dict[str, SignalPin],
        project_b: Project,
        ruleset_b: Optional[Ruleset] = None,
        collector: Optional[ErrorCollector] = None,
        tq_manager: Optional[TQManager] = None,
    ) -> Dict[str, SignalPin]:
        """
        Validate pins against project hardware and new rules:
          - E6001: Pinned slot no longer valid (hardware missing or incompatible).
          - E6002: New rule invalidates pins (new hard rule forbids pinned slot).
        """
        col = collector or ErrorCollector()
        tqs = tq_manager or TQManager()

        # Build hardware coordinate index for project_b
        valid_slots: Set[Tuple[str, str, str, int]] = set()
        cioc_seg_roles: Dict[Tuple[str, str], str] = {}
        cab_areas: Dict[str, Optional[str]] = {}

        for cab in project_b.get("cabinets", []):
            cid = str(cab.get("id", ""))
            cab_areas[cid] = cab.get("area")
            for cioc in cab.get("ciocs", []):
                ciid = str(cioc.get("id", ""))
                cioc_seg_roles[(cid, ciid)] = str(cioc.get("seg_role", "MIXED"))
                for bp in cioc.get("baseplates", []):
                    bpid = str(bp.get("id", ""))
                    cap = int(bp.get("capacity", 12))
                    for s in range(1, cap + 1):
                        valid_slots.add((cid, ciid, bpid, s))

        sigs_b = {str(s.get("tag")): s for s in project_b.get("signals", []) if s.get("tag")}

        for tag, pin in pins.items():
            coord = (pin.cabinet_id, pin.cioc_id, pin.baseplate_id, pin.slot)

            # 1. Hardware validity check (E6001)
            if coord not in valid_slots:
                pin.is_relaxed = True
                pin.relaxation_reason = "Hardware slot no longer exists in project"
                err = make_error(
                    code="E6001",
                    message=f"Pinned slot {pin.cabinet_id}/{pin.cioc_id}/{pin.baseplate_id}/{pin.slot} for signal '{tag}' is no longer valid",
                    phase=Phase.REVISION,
                    severity=Severity.WARNING,
                    context={"tag": tag, "slot": coord},
                    suggestion="Relax pin to allow automatic re-allocation to available hardware",
                )
                col.add(err)
                tqs.create_tq(
                    code="E6001",
                    title=f"Pinned Slot Invalid: {tag}",
                    severity="medium",
                    affected_signals=[tag],
                    evidence=f"Hardware slot {coord} does not exist in Rev B project model",
                    suggestions=["Relax pin and permit relocation"],
                )
                continue

            # 2. Hardware compatibility check (E6001)
            sig_b = sigs_b.get(tag)
            if sig_b:
                sig_area = sig_b.get("area")
                cab_area = cab_areas.get(pin.cabinet_id)
                if sig_area and cab_area and str(sig_area) != str(cab_area):
                    pin.is_relaxed = True
                    pin.relaxation_reason = f"Signal area {sig_area} mismatches cabinet area {cab_area}"
                    err = make_error(
                        code="E6001",
                        message=f"Pinned slot in cabinet '{pin.cabinet_id}' conflicts with signal area '{sig_area}'",
                        phase=Phase.REVISION,
                        severity=Severity.WARNING,
                        context={"tag": tag, "sig_area": sig_area, "cab_area": cab_area},
                        suggestion="Allow signal to move to matching area cabinet",
                    )
                    col.add(err)
                    tqs.create_tq(
                        code="E6001",
                        title=f"Pinned Cabinet Area Mismatch: {tag}",
                        severity="medium",
                        affected_signals=[tag],
                        evidence=f"Signal {tag} requires Area {sig_area}, pinned cabinet is {cab_area}",
                        suggestions=["Permit relocation to area-matched cabinet"],
                    )
                    continue

                # 3. Segregation role check & new rules (E6002)
                seg = sig_b.get("seg")
                cioc_role = cioc_seg_roles.get((pin.cabinet_id, pin.cioc_id), "MIXED")
                if (seg == "A" and cioc_role == "B_ONLY") or (seg == "B" and cioc_role == "A_ONLY"):
                    pin.is_relaxed = True
                    pin.relaxation_reason = f"Segregation {seg} conflicts with CIOC role {cioc_role}"
                    err = make_error(
                        code="E6002",
                        message=f"New rule / segregation role ({cioc_role}) invalidates pinned slot for signal '{tag}' (Seg: {seg})",
                        phase=Phase.REVISION,
                        severity=Severity.WARNING,
                        context={"tag": tag, "seg": seg, "cioc_role": cioc_role},
                        suggestion="Relax pin to allow placement in segregated CIOC",
                    )
                    col.add(err)
                    tqs.create_tq(
                        code="E6002",
                        title=f"New Rule Invalidates Pin: {tag}",
                        severity="medium",
                        affected_signals=[tag],
                        evidence=f"Signal segregation {seg} violates CIOC {pin.cioc_id} segregation role {cioc_role}",
                        suggestions=["Relax pin for segregated re-allocation"],
                    )
                    continue

                # Check active hard rules in ruleset_b for explicit contradictions
                if ruleset_b:
                    for rule in ruleset_b.rules:
                        r_prio = rule.priority or "hard"
                        if r_prio == "hard":
                            r_when_str = str(rule.when) if rule.when else ""
                            applies_to_sig = (not rule.when) or (tag in r_when_str) or ("tag" not in r_when_str)
                            if applies_to_sig:
                                for effect in rule.effects:
                                    if isinstance(effect, ForbidEffect) and isinstance(effect.expr, BinaryOp):
                                        f_field = getattr(effect.expr.left, "field_name", getattr(effect.expr.left, "name", None))
                                        f_val = getattr(effect.expr.right, "value", None)
                                        if effect.expr.op == "==":
                                            if f_field == "slot" and f_val == pin.slot:
                                                pin.is_relaxed = True
                                                pin.relaxation_reason = f"Forbidden by new hard rule {rule.name}"
                                            elif f_field in ("cabinet", "cabinet_id") and f_val == pin.cabinet_id:
                                                pin.is_relaxed = True
                                                pin.relaxation_reason = f"Forbidden by new hard rule {rule.name}"
                                    elif isinstance(effect, RequireEffect) and isinstance(effect.target, BinaryOp):
                                        r_field = getattr(effect.target.left, "field_name", getattr(effect.target.left, "name", None))
                                        r_val = getattr(effect.target.right, "value", None)
                                        if effect.target.op == "==":
                                            if r_field == "slot" and r_val != pin.slot:
                                                pin.is_relaxed = True
                                                pin.relaxation_reason = f"New hard rule requires slot {r_val} (pinned: {pin.slot})"
                                            elif r_field in ("cabinet", "cabinet_id") and r_val != pin.cabinet_id:
                                                pin.is_relaxed = True
                                                pin.relaxation_reason = f"New hard rule requires cabinet {r_val} (pinned: {pin.cabinet_id})"

                                    if pin.is_relaxed:
                                        err = make_error(
                                            code="E6002",
                                            message=f"New rule '{rule.name}' invalidates pinned slot for signal '{tag}': {pin.relaxation_reason}",
                                            rule_id=rule.name,
                                            phase=Phase.REVISION,
                                            severity=Severity.WARNING,
                                            context={"tag": tag, "rule": rule.name, "reason": pin.relaxation_reason},
                                            suggestion="Relax pin to satisfy new rule",
                                        )
                                        col.add(err)
                                        tqs.create_tq(
                                            code="E6002",
                                            title=f"Rule Conflict with Pin: {tag}",
                                            severity="medium",
                                            rules=[rule.name],
                                            affected_signals=[tag],
                                            evidence=f"Rule {rule.name} invalidates pinned location for {tag}",
                                            suggestions=["Relax pin"],
                                        )
                                        break
                            if pin.is_relaxed:
                                break

        return pins

    @staticmethod
    def relax_pins(pins: Dict[str, SignalPin], count_to_relax: int) -> List[SignalPin]:
        """
        Relax pins in strict priority order:
          ungrouped -> soft signals -> group-changed -> safety last
        """
        # Sort active pins by priority ascending (lowest priority first)
        active_pins = [p for p in pins.values() if not p.is_relaxed and p.kind == PinKind.SOFT]
        sorted_pins = sorted(active_pins, key=lambda p: (int(p.priority), p.tag))

        relaxed: List[SignalPin] = []
        for pin in sorted_pins[:count_to_relax]:
            pin.is_relaxed = True
            pin.relaxation_reason = f"Relaxed by priority policy (priority: {pin.priority.name})"
            relaxed.append(pin)

        return relaxed


# =========================================================================
# Excel Reading for Existing Allocation (E6004)
# =========================================================================

def read_allocation_from_excel(
    file_path: Union[str, Path],
    collector: Optional[ErrorCollector] = None,
    tq_manager: Optional[TQManager] = None,
) -> Dict[str, Placement]:
    """
    Load pre-existing Rev A signal allocations from the 'Allocation' sheet of an Excel workbook.
    Raises or records E6004 if the file is missing or has no valid allocation rows.
    """
    col = collector or ErrorCollector()
    tqs = tq_manager or TQManager()

    path = Path(file_path).resolve()
    if not path.exists():
        err = make_error(
            code="E6004",
            message=f"Rev A allocation workbook not found at: {path}",
            phase=Phase.REVISION,
            severity=Severity.ERROR,
            context={"file": str(path)},
            suggestion="Provide a valid Rev A Excel allocation workbook",
        )
        col.add(err)
        tqs.create_tq(
            code="E6004",
            title="Rev A Allocation Missing",
            severity="critical",
            evidence=f"File not found: {path}",
            suggestions=["Check path to Rev A allocation workbook"],
        )
        return {}

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as ex:
        err = make_error(
            code="E6004",
            message=f"Failed to open Rev A workbook {path.name}: {ex}",
            phase=Phase.REVISION,
            severity=Severity.ERROR,
            context={"error": str(ex)},
            suggestion="Ensure workbook is valid .xlsx format",
        )
        col.add(err)
        return {}

    if "Allocation" not in wb.sheetnames:
        err = make_error(
            code="E6004",
            message=f"Workbook {path.name} is missing required 'Allocation' sheet",
            phase=Phase.REVISION,
            severity=Severity.ERROR,
            context={"sheetnames": wb.sheetnames},
            suggestion="Workbook must contain an 'Allocation' sheet",
        )
        col.add(err)
        tqs.create_tq(
            code="E6004",
            title="Rev A Allocation Sheet Missing",
            severity="critical",
            evidence=f"Sheet 'Allocation' not found in {path.name}",
            suggestions=["Provide workbook with 'Allocation' sheet"],
        )
        return {}

    ws = wb["Allocation"]
    placements: Dict[str, Placement] = {}

    # Header row inspection
    headers = [str(cell.value or "").strip().lower() for cell in ws[1]]
    tag_col = next((i for i, h in enumerate(headers, 1) if "tag" in h), 1)
    cab_col = next((i for i, h in enumerate(headers, 1) if "cabinet" in h), 2)
    cioc_col = next((i for i, h in enumerate(headers, 1) if "cioc" in h), 3)
    bp_col = next((i for i, h in enumerate(headers, 1) if "baseplate" in h), 4)
    slot_col = next((i for i, h in enumerate(headers, 1) if "slot" in h), 5)
    charm_col = next((i for i, h in enumerate(headers, 1) if "charm" in h), 6)
    tb_col = next((i for i, h in enumerate(headers, 1) if "tb" in h), 7)

    for row_idx in range(2, ws.max_row + 1):
        tag_val = ws.cell(row=row_idx, column=tag_col).value
        if not tag_val:
            continue
        tag = str(tag_val).strip()

        cab_val = str(ws.cell(row=row_idx, column=cab_col).value or "").strip()
        cioc_val = str(ws.cell(row=row_idx, column=cioc_col).value or "").strip()
        bp_val = str(ws.cell(row=row_idx, column=bp_col).value or "").strip()
        slot_raw = ws.cell(row=row_idx, column=slot_col).value
        try:
            slot_val = int(slot_raw)
        except (ValueError, TypeError):
            continue

        charm_t = str(ws.cell(row=row_idx, column=charm_col).value or "AI").strip()
        tb_t = str(ws.cell(row=row_idx, column=tb_col).value or "STANDARD").strip()

        placements[tag] = Placement(
            tag=tag,
            cabinet_id=cab_val,
            cioc_id=cioc_val,
            baseplate_id=bp_val,
            slot=slot_val,
            charm_type=charm_t,
            tb_type=tb_t,
            applied_rules=["R_RevA_Baseline"],
        )

    if not placements:
        err = make_error(
            code="E6004",
            message=f"Rev A 'Allocation' sheet in {path.name} contains zero valid signal allocations",
            phase=Phase.REVISION,
            severity=Severity.ERROR,
            context={"rows_scanned": ws.max_row - 1},
            suggestion="Populate the Rev A allocation sheet before running revision comparison",
        )
        col.add(err)
        tqs.create_tq(
            code="E6004",
            title="Empty Rev A Allocation",
            severity="critical",
            evidence=f"Zero allocation rows in {path.name}",
            suggestions=["Provide populated Rev A workbook"],
        )

    return placements


# =========================================================================
# Revision Pipeline (10.1 - 10.3)
# =========================================================================

def run_revision_pipeline(
    rev_a_mdb: Union[str, Path, Project],
    rev_b_mdb: Union[str, Path, List[Union[str, Path]]],
    output_path: Optional[Union[str, Path]] = None,
    template_path: Optional[Union[str, Path]] = None,
    rev_a_allocations: Optional[Union[str, Path, Dict[str, Placement]]] = None,
    ruleset_a: Optional[Union[str, Path, Ruleset]] = None,
    ruleset_b: Optional[Union[str, Path, Ruleset]] = None,
    column_mapping: Optional[Dict[str, str]] = None,
    schema: Optional[Schema] = None,
    max_moves: Optional[int] = None,
    solver_backend: str = "cpsat",
    time_limit_seconds: float = 30.0,
    enable_fallback: bool = True,
) -> RevisionResult:
    """
    Execute full revision comparison and move minimization pipeline:
      1. Diff Rev A vs Rev B MDB and rulesets.
      2. Pin unchanged signals to Rev A slots.
      3. Enforce move budget (E6003).
      4. Detect invalid pins (E6001, E6002) and missing Rev A allocation (E6004).
      5. Generate template-preserving output with comprehensive Impact sheet.
    """
    sch = schema or get_base_schema()
    col = ErrorCollector()
    tqs = TQManager()

    # 1. Ingest Project A
    if isinstance(rev_a_mdb, Entity) or (isinstance(rev_a_mdb, dict) and "signals" in rev_a_mdb) or (hasattr(rev_a_mdb, "get") and rev_a_mdb.get("signals") is not None):
        project_a = rev_a_mdb
    else:
        project_a, _, _, _ = read_excel_mdb(
            file_paths=rev_a_mdb,
            column_mapping=column_mapping,
            schema=sch,
            collector=col,
            tq_manager=tqs,
        )

    # 2. Ingest Project B
    if isinstance(rev_b_mdb, Entity) or (isinstance(rev_b_mdb, dict) and "signals" in rev_b_mdb) or (hasattr(rev_b_mdb, "get") and rev_b_mdb.get("signals") is not None):
        project_b = rev_b_mdb
    else:
        project_b, _, _, _ = read_excel_mdb(
            file_paths=rev_b_mdb,
            column_mapping=column_mapping,
            schema=sch,
            collector=col,
            tq_manager=tqs,
        )

    # 3. Load Rev A Placements
    alloc_a: Dict[str, Placement] = {}
    if isinstance(rev_a_allocations, dict):
        alloc_a = rev_a_allocations
    elif rev_a_allocations:
        alloc_a = read_allocation_from_excel(rev_a_allocations, collector=col, tq_manager=tqs)
    elif isinstance(rev_a_mdb, (str, Path)):
        # Try reading Allocation sheet from rev_a_mdb workbook directly
        alloc_a = read_allocation_from_excel(rev_a_mdb, collector=col, tq_manager=tqs)

    if not alloc_a and not col.has_code("E6004"):
        # Rev A allocation missing
        err = make_error(
            code="E6004",
            message="Rev A allocation data is missing or empty",
            phase=Phase.REVISION,
            severity=Severity.ERROR,
            suggestion="Provide Rev A allocation workbook or placement dictionary",
        )
        col.add(err)
        tqs.create_tq(
            code="E6004",
            title="Rev A Allocation Missing",
            severity="critical",
            evidence="No Rev A allocation placements available",
            suggestions=["Provide Rev A allocation data"],
        )

    # 4. Parse Rulesets (if provided)
    rs_a: Optional[Ruleset] = None
    rs_b: Optional[Ruleset] = None

    if isinstance(ruleset_a, Ruleset):
        rs_a = ruleset_a
    elif ruleset_a and Path(ruleset_a).exists():
        from arl.parser import parse_ruleset
        with open(ruleset_a, "r", encoding="utf-8") as fa:
            rs_a = parse_ruleset(fa.read(), collector=col)

    if isinstance(ruleset_b, Ruleset):
        rs_b = ruleset_b
    elif ruleset_b and Path(ruleset_b).exists():
        from arl.parser import parse_ruleset
        with open(ruleset_b, "r", encoding="utf-8") as fb:
            rs_b = parse_ruleset(fb.read(), collector=col)

    # 5. Diffs
    mdb_diff_res = diff_mdb(project_a, project_b)
    ruleset_diff_res = diff_rulesets(rs_a, rs_b)

    # 6. Pin Generation and Validation
    pins = PinManager.create_pins(mdb_diff_res, project_b, alloc_a)
    pins = PinManager.validate_pins(
        pins=pins,
        project_b=project_b,
        ruleset_b=rs_b,
        collector=col,
        tq_manager=tqs,
    )

    # 7. Solve Rev B with Pin Minimization
    if solver_backend == "cpsat":
        from arl.solver_cpsat import CPSATSolver
        solver = CPSATSolver(
            schema=sch,
            tq_manager=tqs,
            collector=col,
            time_limit_seconds=time_limit_seconds,
            enable_fallback=enable_fallback,
        )
    else:
        from arl.solver_greedy import GreedySolver
        solver = GreedySolver(schema=sch, tq_manager=tqs, collector=col)

    alloc_result = solver.solve(project_b, ruleset=rs_b, pins=pins)

    # 8. Compute Moves and Impact
    alloc_diff = diff_allocations(alloc_a, alloc_result.allocations)
    moved_signals: List[MoveRecord] = alloc_diff["moved"]
    move_count = len(moved_signals)
    cabinets_hit = alloc_diff["cabinets_hit"]

    budget_exceeded = False
    if max_moves is not None and move_count > max_moves:
        budget_exceeded = True
        err = make_error(
            code="E6003",
            message=f"Move budget exceeded: required {move_count} moves, but allowed budget is {max_moves}",
            phase=Phase.REVISION,
            severity=Severity.ERROR,
            context={"moves": move_count, "max_moves": max_moves},
            suggestion="Increase move budget or loosen constraints",
        )
        col.add(err)
        tqs.create_tq(
            code="E6003",
            title=f"Move Budget Exceeded: {move_count} > {max_moves}",
            severity="high",
            evidence=f"Allocation required {move_count} moves exceeding budget of {max_moves}",
            suggestions=["Increase max_moves budget", "Pin more signals"],
        )

    # 9. Independent Validation
    validator = Validator()
    val_result = validator.validate(project_b, alloc_result, ruleset=rs_b)

    # 10. Build Impact Summary
    impact_summary = {
        "Status": "Revision Comparison (Rev A vs Rev B)",
        "Signals Added": mdb_diff_res.added_count,
        "Signals Deleted": mdb_diff_res.deleted_count,
        "Signals Moved": move_count,
        "Signals Unchanged": len(alloc_diff["unchanged"]),
        "Total Moves Count": move_count,
        "Move Budget": max_moves if max_moves is not None else "UNLIMITED",
        "Move Budget Exceeded": "YES" if budget_exceeded else "NO",
        "Cabinets Affected Count": len(cabinets_hit),
        "Cabinets Affected": ", ".join(cabinets_hit),
        "Rule Changes Count": ruleset_diff_res.total_changes,
        "summary": {
            "added_count": mdb_diff_res.added_count,
            "deleted_count": mdb_diff_res.deleted_count,
            "moved_count": move_count,
            "unchanged_count": len(alloc_diff["unchanged"]),
            "cabinets_hit": cabinets_hit,
            "move_count": move_count,
            "budget_exceeded": budget_exceeded,
        },
        "rule_changes": [r.model_dump() for r in (
            ruleset_diff_res.added_rules
            + ruleset_diff_res.removed_rules
            + ruleset_diff_res.modified_rules
            + ruleset_diff_res.overridden_rules
        )],
        "moved_signals": [m.model_dump() for m in moved_signals],
    }

    # 11. Excel Export (if output_path is provided)
    out_file_str = None
    if output_path:
        writer = ExcelWriter()
        out_file = writer.write_allocation_workbook(
            output_path=output_path,
            allocation_result=alloc_result,
            project=project_b,
            template_path=template_path,
            ruleset=rs_b,
            validation_result=val_result,
            revision_impact=impact_summary,
        )
        out_file_str = str(out_file)

    is_overall_success = alloc_result.success and val_result.valid and not budget_exceeded and not col.has_code("E6004")

    return RevisionResult(
        success=is_overall_success,
        allocation_result=alloc_result,
        mdb_diff=mdb_diff_res,
        ruleset_diff=ruleset_diff_res,
        moved_signals=moved_signals,
        move_count=move_count,
        max_moves=max_moves,
        budget_exceeded=budget_exceeded,
        cabinets_hit=cabinets_hit,
        tqs=tqs.get_all(),
        impact_summary=impact_summary,
        output_path=out_file_str,
    )
